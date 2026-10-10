#!/usr/bin/env python3
"""Generate document fixtures used by the extractor tests.

Requires: python-docx, openpyxl, python-pptx, reportlab, xlwt, Pillow (with the DejaVu fonts).
Usage: python3 scripts/make-fixtures.py test/fixtures/docs
"""
import json
import os
import sys
import zipfile

out = sys.argv[1] if len(sys.argv) > 1 else "test/fixtures/docs"
os.makedirs(out, exist_ok=True)
p = lambda name: os.path.join(out, name)


def docx():
    from docx import Document

    d = Document()
    d.add_heading("Quarterly Planning Report", 0)
    d.add_paragraph("The marketing budget for the northern region increased by twelve percent.")
    para = d.add_paragraph("Revenue projections ")
    para.add_run("look strong").bold = True
    para.add_run(" for the café expansion.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Region"
    t.cell(0, 1).text = "Headcount"
    t.cell(1, 0).text = "Northwind"
    t.cell(1, 1).text = "42"
    d.sections[0].header.paragraphs[0].text = "Confidential header text"
    d.sections[0].footer.paragraphs[0].text = "Footer page marker"
    d.save(p("report.docx"))


def xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget 2024"
    ws.append(["Department", "Q1", "Q2", "Notes"])
    ws.append(["Engineering", 120000, 135000, "hiring freeze lifted"])
    ws.append(["Marketing", 80000, 92000.5, "conference sponsorship"])
    ws2 = wb.create_sheet("Inventory")
    ws2.append(["SKU", "Item"])
    ws2.append(["ZX-81", "Flux capacitor"])
    ws2["C2"] = True
    wb.save(p("budget.xlsx"))


def pptx():
    from pptx import Presentation

    pr = Presentation()
    for i, (title, body, notes) in enumerate(
        [
            ("Welcome to Orion", "Mission overview and launch timeline", "Remember to thank the volunteers"),
            ("Telemetry", "Downlink rate is 2 Mbps from the probe", ""),
            ("Next Steps", "Schedule the thermal vacuum test", "Ask about the budget"),
        ]
    ):
        s = pr.slides.add_slide(pr.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    pr.save(p("deck.pptx"))


def pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(p("paper.pdf"), pagesize=A4)
    c.setTitle("Photosynthesis paper")
    c.drawString(72, 760, "Photosynthesis converts light energy into chemical energy.")
    c.drawString(72, 740, "Chlorophyll absorbs mostly blue and red wavelengths.")
    c.showPage()
    c.drawString(72, 760, "Second page: the Calvin cycle fixes carbon dioxide.")
    c.showPage()
    c.save()


def odf(name, mimetype, body):
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" office:version="1.2">'
        '<office:automatic-styles><style:style xmlns:style="s" style:name="P1">STYLE NOISE</style:style></office:automatic-styles>'
        f"<office:body>{body}</office:body></office:document-content>"
    )
    with zipfile.ZipFile(p(name), "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), mimetype)
        z.writestr("content.xml", content, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("META-INF/manifest.xml", "<manifest/>", compress_type=zipfile.ZIP_DEFLATED)


def odfs():
    odf(
        "notes.odt",
        "application/vnd.oasis.opendocument.text",
        "<office:text><text:h>Garden journal</text:h><text:p>Planted tomatoes<text:tab/>and basil"
        "<text:s text:c=\"2\"/>today.</text:p><text:p>Rain expected &amp; wind.</text:p></office:text>",
    )
    odf(
        "sheet.ods",
        "application/vnd.oasis.opendocument.spreadsheet",
        "<office:spreadsheet><table:table table:name=\"Sheet1\"><table:table-row>"
        "<table:table-cell><text:p>Apples</text:p></table:table-cell><table:table-cell><text:p>17</text:p></table:table-cell>"
        "</table:table-row></table:table></office:spreadsheet>",
    )
    odf(
        "slides.odp",
        "application/vnd.oasis.opendocument.presentation",
        "<office:presentation><draw:page><draw:frame><draw:text-box><text:p>First slide nebula</text:p></draw:text-box></draw:frame></draw:page>"
        "<draw:page><draw:frame><draw:text-box><text:p>Second slide quasar</text:p></draw:text-box></draw:frame></draw:page></office:presentation>",
    )


def epub():
    with zipfile.ZipFile(p("book.epub"), "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip")
        z.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>',
        )
        z.writestr(
            "OEBPS/content.opf",
            '<package><manifest><item id="c2" href="chap2.xhtml" media-type="application/xhtml+xml"/>'
            '<item id="c1" href="chap1.xhtml" media-type="application/xhtml+xml"/></manifest>'
            '<spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>',
        )
        z.writestr("OEBPS/chap1.xhtml", "<html><head><style>p{}</style></head><body><h1>Chapter One</h1><p>The lighthouse keeper&#8217;s cat.</p></body></html>")
        z.writestr("OEBPS/chap2.xhtml", "<html><body><h1>Chapter Two</h1><p>A storm over the harbour.</p><script>var x='ignored';</script></body></html>")


def rtf():
    with open(p("letter.rtf"), "w") as f:
        f.write(
            r"{\rtf1\ansi\deff0{\fonttbl{\f0 Times New Roman;}}{\colortbl;\red255\green0\blue0;}"
            r"{\info{\title Secret Title}}\f0\fs24 Dear committee,\par We request funding for the "
            r"\b observatory\b0  upgrade.\par Caf\'e9 meeting at noon\emdash bring " + "\\u8364? receipts.\\par}"
        )


def eml():
    with open(p("mail.eml"), "w") as f:
        f.write(
            "From: Ada <ada@example.com>\nTo: Bob <bob@example.com>\n"
            "Subject: =?UTF-8?B?UHJvamVjdCB1cGRhdGUg4pyF?=\nDate: Mon, 1 Jan 2024 10:00:00 +0000\n"
            'MIME-Version: 1.0\nContent-Type: multipart/mixed; boundary="XYZ"\n\n'
            "--XYZ\nContent-Type: multipart/alternative; boundary=\"ALT\"\n\n"
            "--ALT\nContent-Type: text/plain; charset=utf-8\nContent-Transfer-Encoding: quoted-printable\n\n"
            "The deployment is sched=\nuled for Friday. Caf=C3=A9 at 9.\n"
            "--ALT\nContent-Type: text/html; charset=utf-8\n\n<p>HTML version should be skipped</p>\n--ALT--\n"
            "--XYZ\nContent-Type: application/pdf\nContent-Disposition: attachment; filename=x.pdf\nContent-Transfer-Encoding: base64\n\nJVBERi0=\n--XYZ--\n"
        )


def ipynb():
    nb = {
        "cells": [
            {"cell_type": "markdown", "source": ["# Data exploration\n", "Loading the iris dataset"]},
            {"cell_type": "code", "source": ["import pandas as pd\n", "df = pd.read_csv('iris.csv')"], "outputs": []},
        ],
        "metadata": {},
        "nbformat": 4,
        "nbformat_minor": 5,
    }
    with open(p("analysis.ipynb"), "w") as f:
        json.dump(nb, f)


def xls():
    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("Expenses")
    ws.write(0, 0, "Travel")
    ws.write(0, 1, 1234.5)
    ws.write(1, 0, "Lodging in Zürich")
    ws.write(1, 1, 99)
    wb.save(p("legacy.xls"))


def ocr():
    """An image with text, and a PDF of scanned pages (images only), for the OCR helper."""
    from PIL import Image, ImageDraw, ImageFont

    folder = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "test", "fixtures", "ocr")
    os.makedirs(folder, exist_ok=True)
    font = ImageFont.truetype("DejaVuSans.ttf", 48)

    def page(lines, size=(1240, 1754)):
        img = Image.new("RGB", size, "white")
        d = ImageDraw.Draw(img)
        for i, line in enumerate(lines):
            d.text((100, 150 + i * 90), line, fill="black", font=font)
        return img

    page(["Receipt from the Harbour Cafe", "Total paid: 42.50"], (1400, 400)).save(os.path.join(folder, "receipt.png"))
    first = page(["Scanned letter, page one", "Dear tenant, the boiler"])
    second = page(["Page two mentions the zeppelin"])
    first.save(os.path.join(folder, "scanned.pdf"), resolution=150, save_all=True, append_images=[second])


for fn in (docx, xlsx, pptx, pdf, odfs, epub, rtf, eml, ipynb, xls, ocr):
    fn()
print("fixtures written to", out)
